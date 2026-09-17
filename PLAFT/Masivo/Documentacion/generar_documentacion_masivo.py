from pathlib import Path
import csv
import re
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEV = ROOT / "Desarrollo"
OUT = ROOT / "Documentacion"
OUT.mkdir(exist_ok=True)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(8)
            if color:
                run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], header, bold=True, color=(255, 255, 255))
        shade(t.rows[0].cells[i], "1F4E78")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t


def paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def build():
    report = (DEV / "informe_plaft_masivo.html").read_text(encoding="utf-8", errors="ignore")
    dictionary = pd.read_excel(DEV / "diccionario.xlsx")
    importance = pd.read_csv(DEV / "importancia_variables.csv")
    selected = pd.read_csv(DEV / "selected_columns.csv")
    test_vars = pd.read_csv(DEV / "eda_graficos_test" / "variables_modelo.csv")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    for style_name, size, color in [("Title", 22, "1F4E78"), ("Heading 1", 14, "1F4E78"), ("Heading 2", 11, "2F75B5")]:
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.color.rgb = RGBColor.from_string(color)

    header = sec.header.paragraphs[0]
    header.text = "INTERBANK | DATA SCIENCE | PLAFT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Documento metodológico - PN Masivo | Página ")
    add_page_number(footer)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DOCUMENTO METODOLÓGICO\n").bold = True
    title.add_run("MODELO DE PRIORIZACIÓN DE ALERTAS PLAFT\n").bold = True
    title.add_run("Segmento Persona Natural - Masivo").font.size = Pt(16)
    doc.add_paragraph()
    meta = [
        ["Versión", "1.0 - documento inicial generado a partir de evidencia disponible"],
        ["Fecha", date.today().isoformat()],
        ["Área responsable", "Data Science / División Analítica"],
        ["Modelo reportado", "hpo-plaft-pn-masivo-260701-1603-042-a270185f"],
        ["Periodo de test reportado", "202508 - 202605"],
        ["Estado", "Borrador técnico sujeto a reconciliación de artefactos"],
    ]
    table(doc, ["Campo", "Detalle"], meta, [1.8, 4.8])
    paragraph(doc, "Nota de control: este documento consolida los archivos encontrados en PLAFT\\Masivo. No reemplaza la aprobación metodológica ni la validación independiente.")
    doc.add_page_break()

    heading(doc, "1. INTRODUCCIÓN")
    paragraph(doc, "El presente documento describe la metodología de desarrollo, validación e interpretación del Modelo de Priorización de Alertas PLAFT para Personas Naturales del segmento Masivo. El modelo asigna un score entre 0 y 1 para ordenar clientes según prioridad relativa de revisión.")
    paragraph(doc, "El documento se construye siguiendo la estructura del Documento Metodológico PLAFT PJ Minorista v2 y usando como fuentes principales el informe de validación HTML, los notebooks de desarrollo e inferencia, el diccionario de variables y el artefacto XGBoost localizado en la carpeta Masivo.")
    heading(doc, "1.1 Objetivo", 2)
    paragraph(doc, "Apoyar la priorización operativa de clientes y alertas PLAFT mediante un score de riesgo, concentrando la capacidad de análisis en los casos con mayor prioridad relativa.")
    heading(doc, "1.2 Alcance", 2)
    table(doc, ["Elemento", "Definición documentada"], [
        ["Tipología", "Lavado de Activos y Financiamiento del Terrorismo (PLAFT)"],
        ["Segmento", "Persona Natural - Masivo"],
        ["Unidad de análisis", "Cliente"],
        ["Score", "Valor continuo entre 0 y 1, usado para ordenar"],
        ["Población reportada", "13,182,690 clientes en la muestra de perfilamiento"],
    ], [1.8, 4.8])

    heading(doc, "2. CONTEXTO DEL NEGOCIO PLAFT")
    paragraph(doc, "En el segmento Masivo, el modelo busca complementar los mecanismos existentes de monitoreo mediante una priorización basada en señales transaccionales, comportamiento financiero, canales, relaciones con terceros y señales de riesgo PLAFT. El score no constituye por sí mismo una conclusión de operación sospechosa ni sustituye el análisis del especialista.")
    paragraph(doc, "El informe de validación reporta que las señales observadas en comentarios de analistas tienen una cobertura aproximada de 49% por las variables del modelo. Esta cifra debe interpretarse como cobertura de señales identificadas en la muestra analizada, no como una medida de efectividad regulatoria.")

    heading(doc, "3. ARQUITECTURA GENERAL DEL MODELO")
    table(doc, ["Etapa", "Descripción", "Evidencia"], [
        ["Fuentes", "Datos transaccionales, perfil financiero y señales de riesgo PLAFT", "Notebooks de generación de base y querys SQL"],
        ["Ingeniería", "Agregaciones temporales, razones, conteos, montos y señales binarias", "selected_columns.csv / diccionario.xlsx"],
        ["Preprocesamiento", "Preparación de variables y tratamiento requerido por el pipeline", "preprocessing_1.py y notebooks"],
        ["Modelo", "Artefacto serializado XGBoost", "model/model.tar.gz -> xgboost-model"],
        ["Salida", "Score continuo y orden de prioridad", "Informe HTML e inferencia"],
    ], [1.0, 3.5, 2.1])

    heading(doc, "4. CONSTRUCCIÓN DE LA VARIABLE OBJETIVO (TARGET)")
    paragraph(doc, "Los archivos de desarrollo utilizan una columna target binaria y los conjuntos de test incluyen cruces con alertas del periodo 202607. La definición operacional exacta del evento positivo, su ventana de observación y la regla de etiquetado deben quedar ratificadas por el dueño del proceso antes de la aprobación final.")
    paragraph(doc, "Control requerido: documentar explícitamente qué estados de alerta, calificación del analista o resultado posterior conforman target=1 y qué periodo separa la información predictora del evento observado.")

    heading(doc, "5. CONSTRUCCIÓN DEL DATASET")
    paragraph(doc, "La carpeta Masivo contiene notebooks separados para generación de base, EDA, validación final e inferencia. Los nombres de los notebooks indican una partición temporal con meses de entrenamiento, validación y test; la parametrización exacta debe conservarse junto con la ejecución reproducible.")
    table(doc, ["Fuente", "Uso", "Estado"], [
        ["0.genera_base_train.ipynb", "Generación de base completa", "Disponible"],
        ["01.bivariados+EDA_train.ipynb", "EDA y bivariados de train", "Disponible"],
        ["01.bivariados+EDA_test.ipynb", "EDA y bivariados de test", "Disponible"],
        ["02.Validacion_final.ipynb", "Validación del modelo", "Disponible"],
        ["3.VALIDACION_TOTAL...ipynb", "Validación total y señales de analista", "Disponible"],
        ["02.Validacion_RA_por_mes_alertas.ipynb", "Validación en inferencia", "Disponible"],
    ], [2.7, 2.7, 1.2])

    heading(doc, "6. INTEGRIDAD TEMPORAL")
    paragraph(doc, "El informe de validación declara como periodo de test 202508-202605 y fija los cortes de quintiles en 202508. Esta decisión permite comparar perfiles con cortes constantes. La documentación final debe anexar la tabla de meses por partición y la fecha de extracción de cada fuente para demostrar ausencia de fuga temporal.")
    paragraph(doc, "La validación temporal no debe confundirse con una validación de efectividad del proceso PLAFT. Las métricas deben reportarse por mes y junto con el volumen, prevalencia y calidad de las etiquetas.")

    heading(doc, "7. CONSTRUCCIÓN DE VARIABLES PREDICTORAS")
    paragraph(doc, f"El diccionario disponible contiene {len(dictionary)} variables descritas con nombre gerencial, interpretación PLAFT y comparación P1/P5. Las familias observadas incluyen montos y conteos transaccionales, ratios, velocidad y concentración, canales, antigüedad, señales PEP y variables de riesgo.")
    table(doc, ["Familia", "Ejemplos"], [
        ["Montos", "monto_total_dolares, monto_debitos_soles, monto_efectivo_soles"],
        ["Conteos", "n_trx_otros, n_trx_app, n_cuentas_distintas"],
        ["Ratios y z-scores", "ratio_debito_credito, fe_zscore_cargos, fe_ratio_salidas_entradas"],
        ["Temporalidad", "cnt_trx_cargostot_3m, cnt_trx_abonosefect_12m, num_antiguedad"],
        ["Riesgo", "cod_rsg_pep, flg_pep, cod_v13_lugar_op_rsg_12m"],
    ], [1.7, 4.9])

    heading(doc, "8. PREPROCESAMIENTO DE DATOS")
    paragraph(doc, "El preprocesamiento se implementa mediante preprocessing_1.py y los notebooks de desarrollo. Antes de aprobar la versión productiva deben quedar congelados: tratamiento de nulos, tipos, límites de variables, codificación de categóricas, orden de columnas y versión del pipeline. El artefacto de modelo por sí solo no documenta todos esos pasos.")

    heading(doc, "9. SELECCIÓN DE VARIABLES")
    paragraph(doc, f"selected_columns.csv contiene {len(selected)} registros y el archivo importancia_variables.csv contiene {len(importance)} variables con importancia y valor absoluto SHAP. El informe HTML declara 35 variables en el modelo, por lo que estos archivos no pueden tratarse como una única especificación sin reconciliación.")
    top = importance.head(12)
    table(doc, ["Variable", "Importancia", "Valor absoluto SHAP"], [[r["Variable"], f"{r['Importancia']:.6f}", f"{r['Valor Absoluto SHAP']:.6f}"] for _, r in top.iterrows()], [3.2, 1.5, 1.9])
    paragraph(doc, "Regla de control: la lista final de variables debe extraerse del mismo commit o ejecución que produjo el artefacto xgboost-model. Hasta completar ese control, este documento presenta la importancia como evidencia descriptiva y no como contrato productivo.")

    heading(doc, "10. DESARROLLO DEL MODELO")
    paragraph(doc, "El modelo localizado es un XGBoost serializado. La carpeta contiene versiones de notebooks con y sin sufijo de estación de trabajo, por lo que la ejecución aprobada debe identificarse mediante fecha, hash del código, parámetros, semilla, volumen y artefacto resultante.")
    paragraph(doc, "El score se utiliza para ordenar clientes de mayor a menor prioridad. No se documenta una probabilidad calibrada; por ello, el score debe interpretarse como ranking relativo salvo que una calibración independiente sea aprobada.")

    heading(doc, "11. OPTIMIZACIÓN DE HIPERPARÁMETROS")
    paragraph(doc, "El nombre del modelo reportado incluye hpo, lo que indica una búsqueda de hiperparámetros. No se encontró en la carpeta una tabla consolidada de búsqueda con espacio de parámetros, criterio de selección y resultado por configuración. Esa evidencia debe anexarse para cerrar la trazabilidad del entrenamiento.")

    heading(doc, "12. CALIBRACIÓN DEL MODELO")
    paragraph(doc, "No se encontró evidencia suficiente para afirmar que el score esté calibrado como probabilidad de evento. La interpretación aprobada en esta versión es ordinal: un score mayor implica mayor prioridad relativa dentro de la población evaluada.")

    heading(doc, "13. EVALUACIÓN DEL MODELO")
    paragraph(doc, "El informe HTML reporta 13,182,690 clientes, cinco quintiles con cortes fijados en 202508 y un score mediano de 0.0670 en P1 frente a 0.0030 en P5. Estos resultados muestran separación de perfiles en la muestra de perfilamiento; no equivalen por sí solos a AUC, Gini, KS, precisión, recall o lift.")
    table(doc, ["Indicador reportado", "Valor", "Interpretación"], [
        ["Clientes perfilados", "13,182,690", "Volumen de la muestra reportada"],
        ["Variables del modelo", "35", "Declaración del informe HTML; pendiente de reconciliación"],
        ["Mediana score P1", "0.0670", "Quintil de mayor score"],
        ["Mediana score P5", "0.0030", "Quintil de menor score"],
        ["Cobertura de señales de analista", "49%", "Señales observadas cubiertas por variables"],
        ["Periodo de test", "202508-202605", "Declaración del informe HTML"],
    ], [2.3, 1.4, 2.9])
    paragraph(doc, "Pendiente de cierre: anexar métricas discriminatorias y operativas por mes, incluyendo denominadores, prevalencia, intervalos o variabilidad y definición del universo operativo.")

    heading(doc, "14. INTERPRETABILIDAD")
    paragraph(doc, "La interpretabilidad se apoya en importancia global, valores SHAP y perfilamiento por quintiles. P1 representa el mayor score y P5 el menor. El diccionario presenta la lectura de 34 variables y sus relaciones P1/P5; debe completarse con la variable adicional que corresponda al modelo de 35 variables declarado en el informe.")
    paragraph(doc, "El análisis de comentarios de analistas permite contrastar si las señales usadas por los especialistas están representadas por el modelo. Esta evidencia es de apoyo interpretativo y no reemplaza una explicación individual validada para cada cliente.")

    heading(doc, "15. CONCLUSIONES")
    paragraph(doc, "La evidencia disponible permite documentar una versión inicial del Modelo PLAFT PN Masivo como un ranking XGBoost para priorización de revisión. El informe HTML muestra separación entre quintiles, perfiles de score y cobertura parcial de señales de analistas.")
    paragraph(doc, "Antes de declarar el modelo metodológicamente cerrado deben reconciliarse el listado de 35 variables, el diccionario de 34 variables, selected_columns.csv, importancia_variables.csv y variables_modelo.csv; además, deben anexarse las métricas por mes y la definición formal del target.")

    heading(doc, "ANEXO A. DICCIONARIO DE VARIABLES")
    dictionary_rows = []
    for _, row in dictionary.iterrows():
        dictionary_rows.append([row.get("Variable original", ""), row.get("Nombre gerencial", ""), row.get("Descripción entendible", ""), row.get("Interpretación PLAFT", "")])
    table(doc, ["Variable", "Nombre gerencial", "Descripción", "Interpretación PLAFT"], dictionary_rows, [1.6, 1.6, 2.2, 1.5])

    heading(doc, "ANEXO B. MATRIZ DE CONTROL DE EVIDENCIA")
    evidence = [
        ["Documento patrón PJ Minorista v2", "Estructura y secciones", "Disponible", "Usado como formato"],
        ["informe_plaft_masivo.html", "Resultados de perfilamiento y modelo", "Disponible", "Fuente principal de resultados"],
        ["diccionario.xlsx", "Definición e interpretación", "Disponible", "34 variables descritas"],
        ["model.tar.gz", "Artefacto productivo", "Disponible", "Contiene xgboost-model"],
        ["selected_columns.csv", "Lista de variables", "Inconsistente", f"{len(selected)} registros; reconciliar"],
        ["importancia_variables.csv", "Importancia/SHAP", "Inconsistente", f"{len(importance)} registros; reconciliar"],
        ["variables_modelo.csv", "Variables de EDA", "Posible corrida anterior", f"{len(test_vars)} registros; no usar como contrato"],
        ["Métricas AUC/Gini/KS por mes", "Validación cuantitativa", "Faltante", "Generar o anexar desde ejecución aprobada"],
        ["Definición formal del target", "Etiqueta y horizonte", "Faltante", "Ratificar con dueño PLAFT"],
        ["Tabla de partición temporal", "Train/validation/test", "Faltante", "Anexar meses, fechas y volúmenes"],
    ]
    table(doc, ["Artefacto", "Propósito", "Estado", "Acción"], evidence, [2.0, 1.8, 1.2, 1.8])

    heading(doc, "ANEXO C. FUENTES Y REPRODUCIBILIDAD")
    for p in sorted(DEV.glob("*.ipynb")):
        paragraph(doc, p.name)
    paragraph(doc, "Fuentes complementarias: preprocessing_1.py, query_trrain.sql, query_test.sql, query_inferencia.sql, model.tar.gz, informe_plaft_masivo.html, diccionario.xlsx, importancia_variables.csv, selected_columns.csv y los directorios eda_graficos_train, eda_graficos_test, bivariados_graficos_train y bivariados_graficos_test.")

    path = OUT / "Documento_Metodologico_PLAFT_PN_Masivo_v1.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    build()